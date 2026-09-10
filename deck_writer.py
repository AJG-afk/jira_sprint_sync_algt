#!/usr/bin/env python3
"""
deck_writer.py — regenerate the data-driven slides of the Sprint Review deck from the workbook.

Per team 9-slide block, updates ONLY:
  Title slide (+1): "Team – IT OPS", sprint label, Sprint Start/End dates
  Goals slide (+3): swap Goals Completed chart image
  Metrics slide (+4): identity + counts + health tables, swap Sprint Metrics chart
  Kanban slide (+5): kanban table, swap Kanban chart
Narrative slides (PI Commitment, Risks, Demo, Forecast, Feedback) untouched.

Charts matched by TITLE (export order is metrics, goals, kanban — not slide order). Manual
health-table rows (Impediments?, Goals Removed?, Backlog Readiness) left as-is; story counts
written only when present in the workbook row. Also FIXES the "Stations –" title drift.
"""

import re, zipfile
from pptx import Presentation

TEAM_BLOCKS = {
    "NET": {"start": 0,  "title_name": "Networking",  "ident_name": "Networking",      "key": "NET"},
    "STN": {"start": 9,  "title_name": "Stations",    "ident_name": "Stations",        "key": "STN"},
    "WIR": {"start": 18, "title_name": "Windows",     "ident_name": "Windows",         "key": "WIR"},
    "WTE": {"start": 27, "title_name": "Workplace Technology", "ident_name": "Workplace Tech.", "key": "WTE"},
    "JST": {"start": 36, "title_name": "Corp. Efficiency - Jira", "ident_name": "CE – Jira", "key": "JST"},
    "TEL": {"start": 45, "title_name": "Corp. Efficiency - Telephony", "ident_name": "CE – Telephony", "key": "TEL"},
}
VERTICAL = "Enterprise Tech."
SCRUM_LEADER = "Aaron Guenther"
CHART_TITLE_TO_OFFSET = {"sprint metrics": 4, "goals completed": 3, "kanban": 5}


def sprint_number(code):
    if not code:
        return ""
    return str(code).replace(" ", "").split("-")[-1]


def chart_titles(xlsx, sheet_name):
    with zipfile.ZipFile(xlsx) as z:
        wb = z.read('xl/workbook.xml').decode()
        names = re.findall(r'<sheet name="([^"]*)"', wb)
        sheet_no = names.index(sheet_name) + 1
        rels = z.read(f'xl/drawings/_rels/drawing{sheet_no}.xml.rels').decode()
        ridmap = dict(re.findall(r'Id="(rId\d+)"[^>]*chart(\d+)\.xml', rels))
        draw = z.read(f'xl/drawings/drawing{sheet_no}.xml').decode()
        rids = re.findall(r'r:id="(rId\d+)"', draw)
        out = {}
        for pos, rid in enumerate(rids):
            xml = z.read(f'xl/charts/chart{ridmap[rid]}.xml').decode()
            m = re.search(r'<a:t>([^<]*)</a:t>', xml)
            out[pos] = (m.group(1) if m else "").strip()
    return out


def map_charts_to_offsets(xlsx, sheet_name, chart_pngs):
    titles = chart_titles(xlsx, sheet_name)
    result = {}
    for idx, title in titles.items():
        low = title.lower()
        for key, offset in CHART_TITLE_TO_OFFSET.items():
            if key in low and idx in chart_pngs:
                result[offset] = chart_pngs[idx]
    return result


def _set_cell(table, row, col, text):
    cell = table.cell(row, col)
    tf = cell.text_frame
    if tf.paragraphs and tf.paragraphs[0].runs:
        tf.paragraphs[0].runs[0].text = str(text)
        for extra in tf.paragraphs[0].runs[1:]:
            extra.text = ""
    else:
        cell.text = str(text)


def _replace_picture(slide, png_path):
    pics = [sh for sh in slide.shapes if sh.shape_type == 13]
    if not pics:
        return False
    pic = max(pics, key=lambda p: (p.width or 0) * (p.height or 0))
    left, top, width, height = pic.left, pic.top, pic.width, pic.height
    sp = pic._element
    sp.getparent().remove(sp)
    slide.shapes.add_picture(png_path, left, top, width, height)
    return True


def update_title_slide(slide, team_display, sprint_label, start_str, end_str):
    date_pat = re.compile(r'\d{2}/\d{2}/\d{4}')
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        txt = sh.text_frame.text
        if txt.strip() == "Sprint Review":
            continue
        if "IT OPS" in txt:
            for p in sh.text_frame.paragraphs:
                if p.runs:
                    p.runs[0].text = f"{team_display} – IT OPS "
                    for r in p.runs[1:]:
                        r.text = ""
        elif "Sprint Start" in txt or "PI3" in txt:
            runs = [r for p in sh.text_frame.paragraphs for r in p.runs]
            start_i = next((i for i, r in enumerate(runs) if "Sprint Start" in r.text), len(runs))
            if start_i > 0:
                runs[0].text = sprint_label
                for r in runs[1:start_i]:
                    r.text = ""
            date_runs = [r for r in runs if date_pat.search(r.text)]
            if start_str and len(date_runs) >= 1:
                date_runs[0].text = date_pat.sub(start_str, date_runs[0].text)
            if end_str and len(date_runs) >= 2:
                date_runs[1].text = date_pat.sub(end_str, date_runs[1].text)


def update_metrics_tables(slide, ident, counts, health):
    for sh in slide.shapes:
        if not sh.has_table:
            continue
        t = sh.table
        ncol, nrow = len(t.columns), len(t.rows)
        if ncol == 6 and nrow == 2:
            for c, val in enumerate(ident):
                if val is not None:
                    _set_cell(t, 1, c, val)
        elif nrow == 7 and ncol == 2:
            for r, val in enumerate(counts):
                if val is not None:
                    _set_cell(t, r, 1, val)
        elif nrow == 5 and ncol == 2:
            for r, val in health:
                if val is not None:
                    _set_cell(t, r, 1, val)


def update_kanban_table(slide, kanban_vals):
    for sh in slide.shapes:
        if sh.has_table:
            t = sh.table
            for r in range(len(t.rows)):
                if r < len(kanban_vals) and kanban_vals[r] is not None:
                    _set_cell(t, r, 1, kanban_vals[r])
            break


def generate_team(prs, xlsx, team, data, chart_pngs, dates):
    """Apply all data-slide updates for one team block.
    dates = (title_start MM/DD/YYYY, title_end, ident_start 'DD Month YY', ident_end)."""
    blk = TEAM_BLOCKS[team]
    base = blk["start"]
    num = sprint_number(data["sprint"])
    sprint_label = f"{blk['title_name']} – 2026 PI3 Sprint {num}"
    ident_sprint = f"{blk['key']} 26 PI3 – Sprint {num}"
    t_start, t_end, i_start, i_end = dates

    update_title_slide(prs.slides[base + 0], blk["title_name"], sprint_label, t_start, t_end)

    ident = [blk["ident_name"], VERTICAL, SCRUM_LEADER, ident_sprint, i_start, i_end]
    counts = [data["committed"], data["completed"], data["s_committed"], data["s_added"],
              data["s_completed"], data["s_notdone"], data["s_removed"]]
    health = [(0, data["spr_vel_avg"]), (4, data["capacity"])]
    update_metrics_tables(prs.slides[base + 3], ident, counts, health)

    kanban = [data["kanban_pts"], data["kanban_cnt"], data["kanban_avg"],
              data["kanban_avg_prev"], data["spr_vel_avg"], data["spr_vel_avg_prev"]]
    update_kanban_table(prs.slides[base + 4], kanban)

    offset_map = map_charts_to_offsets(xlsx, team, chart_pngs)
    for offset, png in offset_map.items():
        _replace_picture(prs.slides[base + offset - 1], png)
    return {"sprint_label": sprint_label, "charts": len(offset_map)}
